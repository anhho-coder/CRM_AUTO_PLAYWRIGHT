"""Build the quick-reference workbook of the Helpdesk A/B test cases (CRM-11859, Customer Support migration).

Input : test-data/CRM-11859-[Odoo migration] Customer Support module - Test cases/Helpdesk-AB-Test-Cases - Pre-pro.csv
Output: the same folder, "Helpdesk-AB-Test-Cases - quick reference.xlsx"
  - sheet "TC list"    : Category | TC ID | Priority | Role | Type | Quick description (one row per CSV row, CSV order,
                         filter + frozen header). Quick description = the CSV "Feature" column, verbatim.
  - sheet "Statistics" : COUNTIFS tables over "TC list" - Category x Priority, Category x Type, Category x Role.

Rebuild whenever the CSV changes:
  python scripts/build_helpdesk_ab_quick_reference.py [path/to/the.csv]

Nothing is computed by hand: the list is copied cell for cell from the CSV, the statistics are Excel formulas.
"""
import csv
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parent.parent
DEFAULT_CSV = REPO / 'test-data' / 'CRM-11859-[Odoo migration] Customer Support module - Test cases' / 'Helpdesk-AB-Test-Cases - Pre-pro.csv'
OUT_NAME = 'Helpdesk-AB-Test-Cases - quick reference.xlsx'

# Display order for the statistics tables. Values found in the CSV but missing here are appended (sorted), so a new
# value can never be dropped silently.
PRIORITY_ORDER = ['High', 'Medium', 'Low']
TYPE_ORDER = ['happy', 'negative', 'boundary', 'data', 'regression', 'ui', 'role', 'performance']
ROLE_ORDER = ['L1 agent', 'L2 engineer', 'Reviewer (L1 senior)', 'CS manager', 'Administrator', 'Customer (by e-mail)']

# Output columns of "TC list": (header, CSV column)
LIST_COLUMNS = [
    ('Category', 'Area'),
    ('TC ID', 'ID'),
    ('Priority', 'Priority'),
    ('Role', 'Role'),
    ('Type', 'Type'),
    ('Quick description', 'Feature'),
]

FONT = 'Arial'
HEADER_FILL = PatternFill('solid', fgColor='1F4E78')
TOTAL_FILL = PatternFill('solid', fgColor='DDEBF7')
THIN = Side(style='thin', color='BFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def ordered(values, preferred):
    present = list(dict.fromkeys(values))
    head = [v for v in preferred if v in present]
    tail = sorted(v for v in present if v not in preferred)
    return head + tail


def text_cell(ws, row, col, value, **font):
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str) and value.startswith('='):
        cell.data_type = 's'  # a CSV text starting with "=" must stay text, never become a formula
    cell.font = Font(name=FONT, **font)
    return cell


def header_row(ws, row, labels, first_col=1):
    for i, label in enumerate(labels):
        c = text_cell(ws, row, first_col + i, label, bold=True, color='FFFFFF')
        c.fill = HEADER_FILL
        c.border = BOX
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def main():
    csv_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_CSV
    with open(csv_path, encoding='utf-8-sig', newline='') as fh:
        rows = list(csv.DictReader(fh))
    if not rows:
        raise SystemExit(f'{csv_path}: no rows')
    missing = [src for _, src in LIST_COLUMNS if src not in rows[0]]
    if missing:
        raise SystemExit(f'{csv_path}: missing column(s) {missing}')
    ids = [r['ID'] for r in rows]
    dup = [i for i, n in Counter(ids).items() if n > 1]
    if dup:
        raise SystemExit(f'duplicate TC IDs: {dup}')

    categories = list(dict.fromkeys(r['Area'] for r in rows))  # CSV order
    prefix_of = {}
    for r in rows:
        prefix = r['ID'].rsplit('-', 1)[0]  # TC-FLOW-001 -> TC-FLOW
        prefix_of.setdefault(r['Area'], set()).add(prefix)

    wb = Workbook()

    # ---------------- sheet 1: TC list ----------------
    ws = wb.active
    ws.title = 'TC list'
    header_row(ws, 1, [h for h, _ in LIST_COLUMNS])
    for r_i, r in enumerate(rows, start=2):
        for c_i, (_, src) in enumerate(LIST_COLUMNS, start=1):
            text_cell(ws, r_i, c_i, r[src])
    last = len(rows) + 1
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(LIST_COLUMNS))}{last}'
    for col, width in zip('ABCDEF', [34, 14, 10, 22, 13, 100]):
        ws.column_dimensions[col].width = width

    # ---------------- sheet 2: Statistics ----------------
    st = wb.create_sheet('Statistics')
    text_cell(st, 1, 1, 'Helpdesk A/B test cases - quick statistics', bold=True, size=14)
    note = (f"Source: test-data/{csv_path.parent.name}/{csv_path.name} ({len(rows)} rows). "
            "Counts are COUNTIFS formulas over sheet 'TC list'. Quick description = the CSV 'Feature' column. "
            'Rebuild with: python scripts/build_helpdesk_ab_quick_reference.py')
    text_cell(st, 2, 1, note, italic=True, color='595959')
    text_cell(st, 3, 1, 'Rows in TC list', bold=True)
    c = st.cell(row=3, column=2, value=f"=COUNTA('TC list'!$B$2:$B${last})")
    c.font = Font(name=FONT, bold=True)

    rng = {
        'cat': f"'TC list'!$A$2:$A${last}",
        'Priority': f"'TC list'!$C$2:$C${last}",
        'Role': f"'TC list'!$D$2:$D${last}",
        'Type': f"'TC list'!$E$2:$E${last}",
    }
    tables = [
        ('Category x Priority', 'Priority', ordered([r['Priority'] for r in rows], PRIORITY_ORDER)),
        ('Category x Type', 'Type', ordered([r['Type'] for r in rows], TYPE_ORDER)),
        ('Category x Role', 'Role', ordered([r['Role'] for r in rows], ROLE_ORDER)),
    ]

    row = 5
    for title, field, values in tables:
        text_cell(st, row, 1, title, bold=True, size=12)
        row += 1
        head = row
        header_row(st, head, ['Category', 'ID prefix'] + values + ['Total'])
        first_data = head + 1
        for i, cat in enumerate(categories):
            r_ = first_data + i
            text_cell(st, r_, 1, cat).border = BOX
            text_cell(st, r_, 2, ', '.join(sorted(prefix_of[cat]))).border = BOX
            for j, _ in enumerate(values):
                col = 3 + j
                L = get_column_letter(col)
                cell = st.cell(row=r_, column=col, value=f"=COUNTIFS({rng['cat']},$A{r_},{rng[field]},{L}${head})")
                cell.font = Font(name=FONT)
                cell.border = BOX
            tot_col = 3 + len(values)
            cell = st.cell(row=r_, column=tot_col,
                           value=f'=SUM({get_column_letter(3)}{r_}:{get_column_letter(tot_col - 1)}{r_})')
            cell.font = Font(name=FONT, bold=True)
            cell.border = BOX
        total_row = first_data + len(categories)
        text_cell(st, total_row, 1, 'Total', bold=True).fill = TOTAL_FILL
        st.cell(row=total_row, column=2).fill = TOTAL_FILL
        for col in range(3, 3 + len(values) + 1):
            L = get_column_letter(col)
            cell = st.cell(row=total_row, column=col, value=f'=SUM({L}{first_data}:{L}{total_row - 1})')
            cell.font = Font(name=FONT, bold=True)
            cell.fill = TOTAL_FILL
            cell.border = BOX
        row = total_row + 2

    st.column_dimensions['A'].width = 34
    st.column_dimensions['B'].width = 11
    for col in range(3, 3 + max(len(v) for _, _, v in tables) + 1):
        st.column_dimensions[get_column_letter(col)].width = 14

    wb.calculation.fullCalcOnLoad = True  # Excel computes the COUNTIFS when the file is opened
    out = csv_path.parent / OUT_NAME
    wb.save(out)
    print(f'rows={len(rows)} categories={len(categories)} -> {out}')


if __name__ == '__main__':
    main()
